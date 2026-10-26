#!/usr/bin/env python3
"""Retry non-OK rows in an existing OpenCage-geocoded workbook.

Usage:
  python3 -m pip install openpyxl requests
  export OPENCAGE_API_KEY="YOUR_KEY"
  python3 retry_failed_geocodes.py input.xlsx
  python3 retry_failed_geocodes.py input.xlsx output.xlsx

Only rows whose Geocode Status is not OK are considered. Retry Processing Status
makes the run resumable. Confirmed city+country matches update original coordinates;
other results are stored in Candidate columns without overwriting existing coordinates.
"""
import os, re, sys, time, shutil, unicodedata
import requests
from openpyxl import load_workbook

API_URL = "https://api.opencagedata.com/geocode/v1/json"
REQUEST_INTERVAL = 1.05
MAX_RETRIES = 3

COUNTRY_ALIASES = {
    "czech republic": {"czech republic", "czechia"},
    "macau": {"macau", "macao"}, "macao": {"macau", "macao"},
    "hong kong": {"hong kong"},
    "united states": {"united states", "united states of america", "usa"},
    "united kingdom": {"united kingdom", "uk", "great britain"},
    "south korea": {"south korea", "republic of korea"},
    "russia": {"russia", "russian federation"},
    "vietnam": {"vietnam", "viet nam"},
    "iran": {"iran", "islamic republic of iran"},
    "tanzania": {"tanzania", "united republic of tanzania"},
    "moldova": {"moldova", "republic of moldova"},
    "bolivia": {"bolivia", "plurinational state of bolivia"},
    "ivory coast": {"ivory coast", "cote d'ivoire", "côte d'ivoire"},
    "cape verde": {"cape verde", "cabo verde"},
    "swaziland": {"swaziland", "eswatini"},
}

CITY_KEYS = ("city", "town", "village", "municipality", "locality", "suburb",
             "city_district", "borough", "island", "county", "state_district")

def norm(value):
    value = str(value or "").strip().casefold()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(c for c in value if not unicodedata.combining(c))
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()

def country_match(expected, returned):
    expected_norm, returned_norm = norm(expected), norm(returned)
    expected_aliases = {norm(x) for x in COUNTRY_ALIASES.get(expected_norm, {expected_norm})}
    returned_aliases = {norm(x) for x in COUNTRY_ALIASES.get(returned_norm, {returned_norm})}
    return bool(returned_norm and expected_aliases.intersection(returned_aliases))

def city_component(components):
    for key in ("city", "town", "village", "municipality", "locality",
                "suburb", "city_district", "borough", "island", "county",
                "state_district"):
        if components.get(key):
            return str(components[key]).strip()
    return ""

def city_match(expected, components, formatted):
    expected_norm = norm(expected)
    if not expected_norm:
        return False
    candidates = [norm(components.get(k, "")) for k in CITY_KEYS if components.get(k)]
    if expected_norm in candidates:
        return True
    if any(c.startswith(expected_norm + " ") or c.endswith(" " + expected_norm)
           for c in candidates if c):
        return True
    formatted_norm = norm(formatted)
    return formatted_norm == expected_norm or formatted_norm.startswith(expected_norm + " ")

def ensure_col(ws, headers, name):
    if name not in headers:
        headers[name] = ws.max_column + 1
        ws.cell(1, headers[name]).value = name
    return headers[name]

def main():
    if len(sys.argv) not in (2, 3):
        sys.exit("Usage: python3 retry_failed_geocodes.py input.xlsx [output.xlsx]")
    api_key = os.getenv("OPENCAGE_API_KEY")
    if not api_key:
        sys.exit('Set your key first: export OPENCAGE_API_KEY="YOUR_KEY"')

    input_path = os.path.abspath(sys.argv[1])
    output_path = os.path.abspath(sys.argv[2]) if len(sys.argv) == 3 else input_path
    if not os.path.isfile(input_path):
        sys.exit(f"Workbook not found: {input_path}")
    if input_path == output_path:
        backup = input_path + ".bak"
        if not os.path.exists(backup):
            shutil.copy2(input_path, backup)
            print(f"Backup created: {backup}")

    wb = load_workbook(input_path)
    ws = wb.active
    headers = {str(ws.cell(1, c).value).strip(): c
               for c in range(1, ws.max_column + 1) if ws.cell(1, c).value is not None}
    required = ("Country", "City", "Latitude", "Longitude", "Geocode Status")
    missing = [h for h in required if h not in headers]
    if missing:
        sys.exit("Missing required column(s): " + ", ".join(missing))

    country_col, city_col = headers["Country"], headers["City"]
    lat_col, lon_col, status_col = headers["Latitude"], headers["Longitude"], headers["Geocode Status"]
    retry_col = ensure_col(ws, headers, "Retry Processing Status")
    cand_lat = ensure_col(ws, headers, "Candidate Latitude")
    cand_lon = ensure_col(ws, headers, "Candidate Longitude")
    cand_city = ensure_col(ws, headers, "Candidate City")
    cand_country = ensure_col(ws, headers, "Candidate Country")
    cand_formatted = ensure_col(ws, headers, "Candidate Matched Location")
    cand_conf = ensure_col(ws, headers, "Candidate Confidence")
    cand_status = ensure_col(ws, headers, "Candidate Review Status")

    rows = []
    for r in range(2, ws.max_row + 1):
        status = str(ws.cell(r, status_col).value or "").strip().upper()
        retry = str(ws.cell(r, retry_col).value or "").strip().upper()
        if status != "OK" and retry not in {"DONE_MATCH", "DONE_CANDIDATE", "DONE_NO_RESULT"}:
            rows.append(r)

    print(f"Rows to retry: {len(rows)}")
    session, last_request = requests.Session(), 0.0

    for i, r in enumerate(rows, 1):
        country = str(ws.cell(r, country_col).value or "").strip()
        city = str(ws.cell(r, city_col).value or "").strip()
        query = f"{city}, {country}".strip(", ")
        if not city or not country:
            ws.cell(r, retry_col).value = "ERROR_MISSING_CITY_OR_COUNTRY"
            wb.save(output_path)
            continue

        data, error = None, ""
        for attempt in range(MAX_RETRIES):
            wait = REQUEST_INTERVAL - (time.monotonic() - last_request)
            if wait > 0:
                time.sleep(wait)
            try:
                resp = session.get(API_URL, params={
                    "q": query, "key": api_key, "limit": 1, "no_annotations": 1
                }, timeout=30)
                last_request = time.monotonic()
                if resp.status_code == 429 or resp.status_code in (500, 502, 503, 504):
                    error = f"HTTP_{resp.status_code}"
                    time.sleep(2 + attempt * 2)
                    continue
                resp.raise_for_status()
                data = resp.json()
                break
            except requests.RequestException as exc:
                error = str(exc)
                time.sleep(2 + attempt * 2)

        if data is None:
            ws.cell(r, retry_col).value = "ERROR_RETRYABLE"
            ws.cell(r, cand_status).value = error[:300]
            wb.save(output_path)
            print(f"[{i}/{len(rows)}] {query}: API error")
            continue

        api_code = data.get("status", {}).get("code")
        results = data.get("results", [])
        if api_code == 402:
            ws.cell(r, retry_col).value = "ERROR_QUOTA_EXCEEDED"
            wb.save(output_path)
            sys.exit("OpenCage quota exceeded. Progress saved; resume after quota resets.")
        if api_code != 200 or not results:
            ws.cell(r, retry_col).value = "DONE_NO_RESULT"
            ws.cell(r, cand_status).value = f"NO_RESULT_API_STATUS_{api_code}"
            wb.save(output_path)
            print(f"[{i}/{len(rows)}] {query}: no result")
            continue

        result = results[0]
        components = result.get("components", {})
        geometry = result.get("geometry", {})
        formatted = str(result.get("formatted", "")).strip()
        ret_city, ret_country = city_component(components), str(components.get("country", "")).strip()
        lat, lon, confidence = geometry.get("lat"), geometry.get("lng"), result.get("confidence", "")
        cmatch, ctrymatch = city_match(city, components, formatted), country_match(country, ret_country)

        if lat is not None and lon is not None and cmatch and ctrymatch:
            ws.cell(r, lat_col).value = lat
            ws.cell(r, lon_col).value = lon
            ws.cell(r, status_col).value = "OK"
            if "Matched Location" in headers:
                ws.cell(r, headers["Matched Location"]).value = formatted
            ws.cell(r, retry_col).value = "DONE_MATCH"
            ws.cell(r, cand_status).value = "CITY_AND_COUNTRY_MATCH"
            for col in (cand_lat, cand_lon, cand_city, cand_country, cand_formatted, cand_conf):
                ws.cell(r, col).value = None
            print(f"[{i}/{len(rows)}] {query}: matched; coordinates updated")
        else:
            # Preserve original Latitude/Longitude; record OpenCage's probable result.
            ws.cell(r, cand_lat).value = lat
            ws.cell(r, cand_lon).value = lon
            ws.cell(r, cand_city).value = ret_city
            ws.cell(r, cand_country).value = ret_country
            ws.cell(r, cand_formatted).value = formatted
            ws.cell(r, cand_conf).value = confidence
            reasons = []
            if not cmatch: reasons.append("CITY_MISMATCH")
            if not ctrymatch: reasons.append("COUNTRY_MISMATCH")
            ws.cell(r, cand_status).value = ";".join(reasons) or "NO_COORDINATES"
            ws.cell(r, retry_col).value = "DONE_CANDIDATE"
            print(f"[{i}/{len(rows)}] {query}: candidate review ({';'.join(reasons)})")
        wb.save(output_path)

    print(f"Finished. Workbook saved: {output_path}")
    print("Rows already marked OK were not queried. Candidate results do not overwrite original coordinates.")

if __name__ == "__main__":
    main()
